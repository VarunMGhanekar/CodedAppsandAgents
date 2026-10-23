"""Read CSV / XLSX data files into lists of dicts (header row = keys)."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

DATA_SHEET = "Data"


def read_rows(path: str | Path) -> list[dict]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(str(path))
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [dict(r) for r in csv.DictReader(f)]
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        out = []
        for values in rows:
            if values is None or all(v is None for v in values):
                continue
            out.append({h: ("" if v is None else v) for h, v in zip(header, values) if h})
        wb.close()
        return out
    raise ValueError(f"Unsupported file type: {path.name}")


def load_local_data(data_dir: str | Path) -> dict:
    """Load all four inputs from a local folder (used for local runs and tests)."""
    d = Path(data_dir)
    return {
        "drug_master": read_rows(d / "drug_master.csv"),
        "cost": read_rows(d / "pharmacy_acquisition_cost.xlsx"),
        "asp": read_rows(d / "cms_asp_pricing.csv"),
        "commercial": read_rows(d / "commercial_reimbursement.xlsx"),
        "file_versions": {
            "pharmacy_acquisition_cost.xlsx": "mock-v1",
            "cms_asp_pricing.csv": "mock-2026Q4",
            "commercial_reimbursement.xlsx": "mock-v1",
        },
    }

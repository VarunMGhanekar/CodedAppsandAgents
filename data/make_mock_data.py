import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

ARIAL = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", start_color="1F4E78")
WARN = PatternFill("solid", start_color="FFF2CC")

def write_xlsx(path, header, rows, readme, flagged=None, money_cols=()):
    flagged = flagged or {}
    wb = Workbook()
    ws = wb.active; ws.title = "Data"
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(horizontal="center")
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = ARIAL
            if header[c.column-1] in money_cols and isinstance(c.value, (int, float)):
                c.number_format = '$#,##0.00'
    for (r, note) in flagged.items():
        for c in ws[r]: c.fill = WARN
        ws.cell(row=r, column=1).comment = Comment(note, "POC")
    for i, h in enumerate(header, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(h) + 4)
    ws.freeze_panes = "A2"
    rd = wb.create_sheet("README")
    for line in readme: rd.append([line])
    for row in rd.iter_rows():
        for c in row: c.font = ARIAL
    rd["A1"].font = Font(name="Arial", size=11, bold=True)
    rd.column_dimensions["A"].width = 110
    wb.save(path)

# ---------------- Drug master ----------------
with open("drug_master.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["brand", "jCode", "biosimilar", "qCode", "active"])
    for b, q in [("Fulphila", "Q5108"), ("Udenyca", "Q5111"),
                 ("Ziextenzo", "Q5120"), ("Nyvepria", "Q5122")]:
        w.writerow(["Neulasta", "J2506", b, q, "true"])

# ---------------- Acquisition cost ----------------
cost_header = ["hcpcs", "ndc", "product", "channel", "cost_per_unit", "effective_date"]
cost_rows = [
    ["Q5108", "67457-0833-06", "Fulphila", "GPO", 52.00, "2026-10-01"],
    ["Q5108", "67457-0833-06", "Fulphila", "340B", 31.20, "2026-10-01"],
    ["Q5111", "70114-0101-01", "Udenyca", "GPO", 48.50, "2026-10-01"],
    ["Q5111", "70114-0101-01", "Udenyca", "340B", 29.10, "2026-10-01"],
    ["Q5120", "61314-0866-01", "Ziextenzo", "GPO", 50.00, "2026-10-01"],
    # Ziextenzo 340B intentionally missing  -> "Cost missing"
    ["Q5122", "0069-0324-01", "Nyvepria", "GPO", 45.75, "2026-10-01"],
    ["Q5122", "0069-0324-01", "Nyvepria", "340B", 27.45, "2026-10-01"],
    ["Q5130", "00000-0000-00", "Fylnetra", "GPO", -10.00, "2026-10-01"],   # invalid: negative
    ["Q51X1", "00000-0000-00", "Unknown", "GPO", 40.00, "2026-10-01"],     # invalid: bad code
]
write_xlsx("pharmacy_acquisition_cost.xlsx", cost_header, cost_rows,
    ["MOCK DATA - Pharmacy acquisition cost (per HCPCS billing unit, 0.5 mg pegfilgrastim)",
     "All prices are invented for the POC. NDCs are placeholders. Verify HCPCS codes before demo.",
     "Edit only the 'Data' sheet. Row 1 is the header; one row per product per channel.",
     "Example row: Q5108 | 67457-0833-06 | Fulphila | GPO | 52.00 | 2026-10-01",
     "",
     "Deliberate test cases:",
     " - Ziextenzo has no 340B row  -> Cost missing warning for 340B requests",
     " - Row 9 (Fylnetra) negative price -> Invalid data, row skipped",
     " - Row 10 bad HCPCS code 'Q51X1'   -> Invalid data, row skipped"],
    flagged={9: "Deliberate invalid data: negative price", 10: "Deliberate invalid data: bad HCPCS"},
    money_cols=("cost_per_unit",))

# ---------------- CMS ASP ----------------
with open("cms_asp_pricing.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["hcpcs", "description", "payment_limit", "quarter"])
    for row in [["J2506", "Inj pegfilgrastim excl bio 0.5 mg", "212.00", "2026Q4"],
                ["Q5108", "Inj Fulphila 0.5 mg", "58.40", "2026Q4"],
                ["Q5111", "Inj Udenyca 0.5 mg", "61.25", "2026Q4"],
                ["Q5120", "Inj Ziextenzo 0.5 mg", "55.10", "2026Q4"],
                ["Q5122", "Inj Nyvepria 0.5 mg", "57.80", "2026Q4"]]:
        w.writerow(row)

# ---------------- Commercial ----------------
com_header = ["payer", "region", "hcpcs", "reimbursement_per_unit"]
rates = {
    ("Aetna", "South"): {"Q5108": 70.00, "Q5111": 66.00, "Q5120": 64.00, "Q5122": 63.00},
    ("Aetna", "West"):  {"Q5108": 68.00, "Q5111": 69.00, "Q5120": 62.00},  # Nyvepria missing
    ("UnitedHealthcare", "South"): {"Q5108": 69.00, "Q5111": 66.50, "Q5120": 67.00, "Q5122": 63.75},
    ("UnitedHealthcare", "West"):  {"Q5108": 65.00, "Q5111": 64.00, "Q5120": 63.00, "Q5122": 60.00},
}
com_rows = [[p, r, code, v] for (p, r), m in rates.items() for code, v in m.items()]
write_xlsx("commercial_reimbursement.xlsx", com_header, com_rows,
    ["MOCK DATA - Commercial payer reimbursement (per HCPCS billing unit)",
     "All rates are invented for the POC. Edit only the 'Data' sheet.",
     "Example row: Aetna | South | Q5108 | 70.00",
     "",
     "Deliberate test cases:",
     " - Aetna / West has no Nyvepria (Q5122) rate -> Reimbursement rate missing warning",
     " - UnitedHealthcare / South GPO: Udenyca and Nyvepria both have 18.00 margin -> tie-break picks Nyvepria (lower cost)"],
    money_cols=("reimbursement_per_unit",))
print("done")
